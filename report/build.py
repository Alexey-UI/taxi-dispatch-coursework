"""Build editable coursework and slides. Dependencies: python-docx, python-pptx, Pillow.

Run from anywhere: python report/build.py. PDFs are exported separately by LibreOffice.
No personal details or secrets are read by this generator.
"""
from pathlib import Path
import json
import re
import shutil
import copy
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt as SlidePt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
ASSETS.mkdir(exist_ok=True)
OUT = ROOT / 'output'
OUT.mkdir(exist_ok=True)
FONT = '/System/Library/Fonts/Supplemental/Arial.ttf'
if not Path(FONT).exists():
    FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'

def font(size):
    return ImageFont.truetype(FONT, size)

def diagram(filename, title, boxes, arrows, size=(1600, 850)):
    img = Image.new('RGB', size, '#fbfcf8')
    d = ImageDraw.Draw(img)
    d.text((50, 30), title, fill='#1d2a22', font=font(38))
    for x1,y1,x2,y2 in arrows:
        d.line((x1,y1,x2,y2), fill='#84946b', width=5)
        import math
        angle = math.atan2(y2-y1,x2-x1)
        points = [(x2,y2),(x2-20*math.cos(angle-.45),y2-20*math.sin(angle-.45)),(x2-20*math.cos(angle+.45),y2-20*math.sin(angle+.45))]
        d.polygon(points, fill='#84946b')
    for x,y,w,h,lines in boxes:
        d.rounded_rectangle((x,y,x+w,y+h), radius=20, fill='#e8efdd', outline='#879c6c', width=3)
        for i,line in enumerate(lines):
            f=font(26 if i == 0 else 21)
            bbox=d.textbbox((0,0),line,font=f)
            d.text((x+(w-bbox[2])/2,y+24+i*36),line,font=f,fill='#223320')
    img.save(ASSETS / filename)

diagram('architecture.png','Компоненты и сетевые связи',[
 (50,130,330,120,['Браузер','Роли и публичная карта']),
 (520,130,340,120,['Nginx / web','Traefik перед web в K3s']),
 (1010,130,480,120,['Go API × 2','HTTP, роли, заказы']),
 (60,420,330,120,['Node.js simulator','Сценарии и движение']),
 (570,420,340,120,['PostgreSQL','Заказы, водители, events']),
 (1080,420,390,120,['Redis','Сигнал новой работы']),
 (570,660,340,120,['Go dispatcher × 2','Поиск, тайм-аут, назначение']),
],[(380,190,520,190),(860,190,1010,190),(1250,250,740,420),(1250,250,1275,420),(390,450,1010,210),(740,660,740,540),(910,720,1250,540)])
diagram('data.png','Сущности и ограничения',[
 (65,200,380,220,['drivers','id, class, x, y, status','available / offered / busy']),
 (580,200,450,300,['orders','id, owner_id, class, route','status, driver_id, expires','UNIQUE: активный driver_id']),
 (1160,200,360,220,['attempts','order_id + driver_id','Окно повторной попытки']),
 (580,610,450,150,['events','service, kind, order_id, time'])
],[(445,300,580,300),(1030,300,1160,300),(800,500,800,610)])
diagram('lifecycle.png','Состояния заказа и повторное назначение',[
 (50,180,300,120,['pending','Ожидает кандидата']),
 (480,180,300,120,['offered','12 секунд на ответ']),
 (930,180,300,120,['accepted','Водитель едет к подаче']),
 (930,460,300,120,['arrived','Водитель прибыл']),
 (480,460,300,120,['in_progress','Поездка']),
 (50,460,300,120,['completed','Завершение']),
 (1190,650,340,120,['cancelled','До начала поездки']),
 (50,650,730,120,['Отказ или тайм-аут: offered → pending','Водитель освобождается, поиск повторяется']),
],[(350,240,480,240),(780,240,930,240),(1080,300,1080,460),(930,520,780,520),(480,520,350,520)])
diagram('pipeline.png','Проверка и доставка GitHub Actions',[
 (50,160,420,140,['Git push / pull request','Conventional Commits']),
 (580,160,420,140,['Проверки Go и JavaScript','vet, test -race, Node tests']),
 (1110,160,420,140,['Compose-интеграция','Два worker; отказ Redis']),
 (1110,490,420,150,['Временный Kubernetes','k3d → Helm → HTTP smoke']),
 (580,490,420,150,['Публикация GHCR','amd64 + arm64; тег SHA']),
 (50,490,420,150,['Облачная ВМ — позже','Отдельный ручной выпуск']),
],[(470,230,580,230),(1000,230,1110,230),(1320,300,1320,490),(1110,565,1000,565),(580,565,470,565)])

screenshot = ROOT.parent / 'artifacts/screenshots/desktop.png'
if screenshot.exists():
    shutil.copyfile(screenshot, ASSETS / 'desktop.png')

text = (ROOT / 'report.md').read_text()
pages = json.loads((ROOT/'pages.json').read_text()) if (ROOT/'pages.json').exists() else {}
doc = Document()
section = doc.sections[0]
section.page_width,section.page_height=Cm(21),Cm(29.7)
section.left_margin,section.right_margin=Cm(3),Cm(1.5)
section.top_margin,section.bottom_margin=Cm(2),Cm(2)
normal=doc.styles['Normal']
normal.font.name='Times New Roman'; normal.font.size=Pt(14)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.first_line_indent=Cm(1.25)
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
for name in ['Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(14);style.font.bold=True
    style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    style.paragraph_format.first_line_indent=Cm(0 if name=='Heading 1' else 1.25)
    style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(12)
    style.paragraph_format.line_spacing=1
    style.paragraph_format.keep_with_next=True

def p(text='',center=False,bold=False,size=None):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if center:para.paragraph_format.first_line_indent=Cm(0)
    run=para.add_run(text);run.bold=bold
    if size:run.font.size=Pt(size)
    return para

for line in ['МИНОБРНАУКИ РОССИИ','Федеральное государственное бюджетное образовательное учреждение высшего образования','«МИРЭА — Российский технологический университет»','РТУ МИРЭА','Институт информационных технологий','Кафедра инструментального и прикладного программного обеспечения']:
    p(line,True,size=12)
p();p('КУРСОВАЯ РАБОТА',True,True)
p('по дисциплине «Настройка и администрирование сервисного программного обеспечения»',True)
p('Направление: 09.03.04 «Программная инженерия»',True,size=12)
p('Профиль: Разработка программных продуктов и проектирование информационных систем',True,size=12)
p();p('Тема: «Настройка и администрирование системы автоматической обработки и распределения заказов для службы такси»',True,True)
p();p('Студент: __________________________________');p('Группа: ___________________________________');p('Руководитель: ______________________________')
p('Представлена к защите: ______________________',size=12);p('Допущена к защите: _________________________',size=12);p('Оценка и подписи комиссии: __________________',size=12)
p('Москва — 2026',True)
doc.add_page_break()
for line in ['МИНОБРНАУКИ РОССИИ','«МИРЭА — Российский технологический университет»','Институт информационных технологий','Кафедра инструментального и прикладного программного обеспечения']:
    p(line,True,size=12)
p('ЗАДАНИЕ НА КУРСОВУЮ РАБОТУ',True,True)
p('По дисциплине: Настройка и администрирование сервисного программного обеспечения',size=12)
p('Студент: _____________________  Группа: ______________')
p('Руководитель: ______________________________________')
p('Срок представления: _________________________________')
p('Тема: Настройка и администрирование системы автоматической обработки и распределения заказов для службы такси.')
p('Исходные данные: Go, JavaScript, PostgreSQL, Redis, Docker, Docker Compose, Kubernetes, Helm; наличие межконтейнерной сети, кластерной архитектуры и CI/CD.')
p('Перечень вопросов: анализ предметной области; обоснование технологий; создание пяти и более контейнеров; кластерная архитектура; CI/CD; деплой; сетевое взаимодействие; презентация и документация.')
p('Руководителем проведён инструктаж по технике безопасности, противопожарной технике и правилам внутреннего распорядка.',size=12)
p('Зав. кафедрой: __________________  Дата: __________',size=12)
p('Задание выдал: __________________  Дата: __________',size=12)
p('Задание получил: ________________  Дата: __________',size=12)
section=doc.add_section(WD_SECTION_START.NEW_PAGE)
section.footer.is_linked_to_previous=False
foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

headings=[line.lstrip('# ').strip() for line in text.splitlines() if line.startswith(('# ','## ')) and line[2:]!='АННОТАЦИЯ']
def contents():
    doc.add_page_break();p('СОДЕРЖАНИЕ',True,True)
    for h in headings:
        para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1.5
        para.paragraph_format.tab_stops.add_tab_stop(Cm(16.2),__import__('docx').enum.text.WD_TAB_ALIGNMENT.RIGHT,__import__('docx').enum.text.WD_TAB_LEADER.DOTS)
        para.paragraph_format.first_line_indent=Cm(.5 if re.match(r'^\d+\.\d+',h) else 0)
        run=para.add_run(h+'\t'+str(pages.get(h,'—')));run.font.size=Pt(14)

lines=text.splitlines();i=0;first=True;toc=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        title=line[2:]
        if not first and not toc:contents();toc=True
        h=doc.add_paragraph(title,'Heading 1');h.alignment=WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.page_break_before=not first
        first=False;i+=1;continue
    if line.startswith('## '):
        previous=next((x.strip() for x in reversed(lines[:i]) if x.strip()),'')
        h=doc.add_paragraph(line[3:],'Heading 2');h.paragraph_format.page_break_before=not previous.startswith('# ')
        i+=1;continue
    if line.startswith('```'):
        i+=1;code=[]
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1
        run=para.add_run('\n'.join(code));run.font.name='Courier New';run.font.size=Pt(9)
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            row=[cell.strip() for cell in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',c or '-') for c in row):rows.append(row)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
        for n,row in enumerate(rows):
            cells=table.add_row().cells
            for cell,value in zip(cells,row):
                para=cell.paragraphs[0];para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1;para.paragraph_format.space_after=Pt(4)
                para.paragraph_format.keep_with_next=n<len(rows)-1
                run=para.add_run(value);run.font.size=Pt(11);run.bold=n==0
            if n==0:
                repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        continue
    match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
    if match:
        caption,path=match.groups();para=doc.add_paragraph();para.paragraph_format.first_line_indent=Cm(0);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(ROOT/path),width=Cm(16.2))
        para=p(caption,True);para.paragraph_format.line_spacing=1;para.runs[0].font.size=Pt(14)
        i+=1;continue
    para=p(line)
    if line.startswith(('Таблица ', 'Листинг ')):
        para.paragraph_format.keep_with_next=True
        para.paragraph_format.first_line_indent=Cm(0)
        para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    i+=1

doc.core_properties.title='Настройка и администрирование системы автоматической обработки и распределения заказов для службы такси'
doc.core_properties.author=''
doc.save(OUT/'Реферат_Тактика.docx')

# Separate editable forms required by the methodological package.
elements=list(doc.element.body)
first_break=next(i for i,e in enumerate(elements) if e.xpath('.//w:br[@w:type="page"]'))
section_break=next(i for i,e in enumerate(elements) if e.xpath('.//w:sectPr'))
for filename,selected in [('Титульный_лист',elements[:first_break]),('Лист_задания',elements[first_break+1:section_break])]:
    form=Document(OUT/'Реферат_Тактика.docx')
    body=form.element.body
    sect=copy.deepcopy(body.sectPr)
    for ref in list(sect):
        if ref.tag in [qn('w:footerReference'),qn('w:headerReference')]:sect.remove(ref)
    for child in list(body):body.remove(child)
    for child in selected:body.append(copy.deepcopy(child))
    body.append(sect)
    form.save(OUT/(filename+'.docx'))

application=Document()
app_section=application.sections[0]
app_section.left_margin,app_section.right_margin=Cm(3),Cm(1.5)
app_section.top_margin,app_section.bottom_margin=Cm(2),Cm(2)
app_section.page_width,app_section.page_height=Cm(21),Cm(29.7)
style=application.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(14)
style.paragraph_format.line_spacing=1.5
for line in ['Заведующему кафедрой инструментального','и прикладного программного обеспечения','Института информационных технологий','________________________________','От студента ____________________','Группа __________  Курс ________','Контакт: _______________________']:
    para=application.add_paragraph(line);para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
para=application.add_paragraph('ЗАЯВЛЕНИЕ');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.runs[0].bold=True
para=application.add_paragraph('Прошу утвердить тему курсовой работы по дисциплине «Настройка и администрирование сервисного программного обеспечения» образовательной программы бакалавриата 09.03.04 «Программная инженерия»: «Настройка и администрирование системы автоматической обработки и распределения заказов для службы такси».');para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;para.paragraph_format.first_line_indent=Cm(1.25)
for line in ['Приложение: проект листа задания на курсовую работу.','Подпись студента: ______________  Дата: __________','Руководитель: ________________________________','Подпись руководителя: __________  Дата: _________','Утверждаю: ___________________  Дата: _________']:
    application.add_paragraph(line)
application.save(OUT/'Заявление_на_тему.docx')

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def slide(title,body=None,image=None,note=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('F5F6F0')
    shape=s.shapes.add_textbox(Inches(.6), Inches(.45), Inches(12), Inches(.85));tf=shape.text_frame;tf.text=title
    tf.paragraphs[0].font.size=SlidePt(29);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=RGBColor.from_string('213321')
    if image:
        img=Image.open(ASSETS/image);ratio=img.width/img.height
        w=min(12,5.55*ratio);h=w/ratio
        s.shapes.add_picture(str(ASSETS/image), Inches((13.333-w)/2), Inches(1.4),width=Inches(w),height=Inches(h))
    if body:
        box=s.shapes.add_textbox(Inches(.8),Inches(1.8),Inches(11.7),Inches(4.9));box.text_frame.word_wrap=True
        for j,line in enumerate(body):
            para=box.text_frame.paragraphs[0] if j==0 else box.text_frame.add_paragraph();para.text=line;para.font.size=SlidePt(23);para.space_after=SlidePt(22);para.font.color.rgb=RGBColor.from_string('33442D')
    footer=s.shapes.add_textbox(Inches(.6),Inches(7.12),Inches(12),Inches(.25));footer.text_frame.text=f'РТУ МИРЭА · Тактика · 2026                                                               {len(prs.slides):02}'
    footer.text_frame.paragraphs[0].font.size=SlidePt(10)
    s.notes_slide.notes_text_frame.text=note

slide('Настройка и администрирование системы такси',[
 'Система автоматической обработки и распределения заказов «Тактика»',
 'Дисциплина: Настройка и администрирование сервисного ПО',
 'Студент: __________________   Группа: __________',
 'Руководитель: __________________',
 'МИНОБРНАУКИ РОССИИ · РТУ МИРЭА · ИТ · ИиППО · Москва, 2026'
],note='Представить тему. До сдачи заполнить данные автора и руководителя.')
slide('Цель и задачи',[
 'Разработать и настроить воспроизводимый стенд распределения заказов.',
 'Выбор ближайшего водителя, отказы, тайм-ауты и целостность состояния.',
 'Контейнеризация, кластер, сеть и автоматизированная доставка.',
 'Проверка штатных сценариев, восстановления и работающие инструкции.'
],note='Около 25 секунд: акцент на настройке и эксплуатации, бизнес-функции ограничены.')
slide('Анализ предметной области',[
 'Такси-Мастер: комплексные роли и правила распределения.',
 'Яндекс Про: автоматический подбор с учётом требований и приоритета.',
 'Учебная модель: доступность + класс + расстояние; явные переходы.',
 'Граница: синтетический город, без оплаты и дорожного графа.'
],note='Около 30 секунд. Не утверждать знание закрытого стека коммерческих аналогов.')
slide('Технологический комплекс',[
 'Go: HTTP API и отдельный dispatcher. JavaScript: интерфейс.',
 'Node.js: воспроизводимые сценарии и движение автомобилей.',
 'PostgreSQL: источник истины. Redis: ускоряющее уведомление.',
 'Nginx · Docker Compose · K3s/k3d · Helm · GitHub Actions.'
],note='Около 30 секунд. Объяснить осмысленную роль каждого компонента.')
slide('Структура и взаимодействие',image='architecture.png',note='Около 35 секунд. Внешний вход только web; общий SQL и ограниченные Redis-сигналы.')
slide('Распределение и жизненный цикл',image='lifecycle.png',note='Около 35 секунд. Блокировки и уникальный индекс запрещают двойное назначение; отказ возвращает поиск.')
slide('Демонстрация работающей системы',image='desktop.png',note='Около 50 секунд. Показать карту, сценарии и события. При живом показе проследить один заказ.')
slide('CI/CD: реальная проверка доставки',image='pipeline.png',note='Около 35 секунд. Успешный run 40655602805; облачная ВМ пока не развёрнута.')
slide('Результаты и границы',[
 'Шесть компонентов Compose; кластер из трёх локальных узлов.',
 'Проверены два dispatcher, отказы, тайм-ауты и работа без Redis.',
 'Подтверждены восстановление API-подов и SQL-дампа.',
 'CI/CD выполнен; образы опубликованы для amd64 и arm64.',
 'Ограничения: одиночная БД и хост. Домен/HTTPS подготовлены к облаку.'
],note='Около 25 секунд. Ссылка на исходники: github.com/Alexey-UI/taxi-dispatch-coursework.')
prs.save(OUT/'Презентация_Тактика.pptx')
print('Built DOCX and PPTX:',OUT)
